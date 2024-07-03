import os
from openpyxl import load_workbook, Workbook

# Ensure the Desktop path
desktop_path = os.path.join(os.path.expanduser("~"), "Desktop")
file_path = os.path.join(desktop_path, 'GST Calc.xlsx')

try:
    # Load the workbook and sheet
    wb = load_workbook(filename=file_path)
except FileNotFoundError:
    wb = Workbook()
    ws = wb.active
else:
    wb = load_workbook(filename=file_path)
    ws = wb.active

# Add data to the sheet
row_num = ws.max_row + 1
ws.cell(row=row_num, column=1).value = '23 Feb'
ws.cell(row=row_num, column=2).value = 'BioMed'
ws.cell(row=row_num, column=3).value = 198.24
ws.cell(row=row_num, column=4).value = 8.85
ws.cell(row=row_num, column=5).value = 12.39

# Save the workbook
wb.save(file_path)

print(f"Data added to Excel file 'GST Calc.xlsx'!")